"""
Parlem — backend de traducción con IA GRATUITA (valenciano <-> español)
==============================================================

Motor: Gemini (API gratuita de Google AI Studio). Entiende contexto, registro
y matices, y además "ve" imágenes y PDF, así que no hace falta Tesseract.

Endpoints:
  GET  /                  -> healthcheck
  POST /chat              -> JSON: {"mensaje": "...", "historial": [{"role": "user|assistant", "content": "..."}],
                                     "direccion": "auto|val-spa|spa-val"}
                             Chatbot conversacional: traduce, explica gramática, etc.
  POST /traducir-texto    -> JSON: {"texto": "...", "direccion": "auto|val-spa|spa-val"}
                             Traducción directa (sin conversación).
  POST /traducir-archivo  -> multipart: file=<docx|pptx|pdf>, direccion=..
                             Devuelve el documento traducido.
  POST /traducir-imagen   -> multipart: file=<png|jpg|webp|gif>, direccion=..
                             Devuelve {"texto_original": "...", "traduccion": "..."}

Variables de entorno:
  GEMINI_API_KEY  (obligatoria; gratis en https://aistudio.google.com/apikey)
  GEMINI_MODEL    (opcional, por defecto gemini-3.5-flash)
  PORT               (opcional, por defecto 5000)
"""

import base64
import io
import json
import logging
import os
import re
from xml.sax.saxutils import escape

from google import genai
from google.genai import types
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("parlem-backend")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB por archivo
CORS(app)

MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.5-flash")
client = genai.Client()  # lee GEMINI_API_KEY del entorno

DIRECCIONES = {"auto", "val-spa", "spa-val"}

# ---------------------------------------------------------------------------
# Prompts
# ---------------------------------------------------------------------------
BASE_RULES = """\
Eres Parlem, un traductor experto entre VALENCIANO y ESPAÑOL (castellano).
- Escribe el valenciano según la normativa de l'Acadèmia Valenciana de la Llengua (AVL):
  formas valencianas (p. ej. "cantar", "xicotet", "huit", "este/eixe/aquell", "açò", "vosté"),
  no catalán central.
- Conserva el tono (formal/informal), la puntuación, los saltos de línea, los nombres propios,
  cifras, URLs y el formato general del texto original.
- No añadas explicaciones, notas ni comentarios a una traducción, salvo que el usuario los pida.
"""


def direction_rule(direccion):
    if direccion == "val-spa":
        return "Dirección: traduce de valenciano a español."
    if direccion == "spa-val":
        return "Dirección: traduce de español a valenciano."
    return (
        "Dirección: detecta el idioma del texto. Si está en valenciano (o catalán), "
        "tradúcelo a español; si está en español, tradúcelo a valenciano."
    )


def clean_direction(value):
    return value if value in DIRECCIONES else "auto"


CHAT_SYSTEM = (
    BASE_RULES
    + """
Eres también un asistente conversacional amable. Reglas del chat:
- Si el usuario te pega un texto sin más instrucciones, tradúcelo y devuelve solo la traducción.
- Si pregunta sobre gramática, vocabulario, diferencias valenciano/catalán/español, ortografía,
  conjugaciones, etc., responde con claridad y ejemplos, en español salvo que pida otro idioma.
- Si te piden algo ajeno a la lengua (no es tu función), indica amablemente que tu especialidad
  es el valenciano y el español, y ofrece ayuda en eso.
- Si te piden traducir a otros idiomas (p. ej. inglés), explica que de momento solo trabajas
  con valenciano y español.
"""
)

# ---------------------------------------------------------------------------
# Utilidades de llamada a la IA
# ---------------------------------------------------------------------------


def _to_parts(content):
    """Convierte contenido (texto o bloques image/document en base64) a partes de Gemini."""
    if isinstance(content, str):
        return [types.Part.from_text(text=content)]
    parts = []
    for block in content:
        if block["type"] == "text":
            parts.append(types.Part.from_text(text=block["text"]))
        else:  # image / document
            src = block["source"]
            parts.append(
                types.Part.from_bytes(data=base64.b64decode(src["data"]), mime_type=src["media_type"])
            )
    return parts


FALLBACK_MODELS = [
    m.strip()
    for m in os.environ.get("GEMINI_FALLBACKS", "gemini-3.1-flash-lite,gemini-3-flash-preview").split(",")
    if m.strip()
]


def ask_ai(system, messages, max_tokens=None):
    """Llama a Gemini. Si el modelo está saturado (503) reintenta con espera;
    si sigue fallando, o no existe / agotó su cuota, prueba el siguiente modelo."""
    import time

    contents = [
        types.Content(role="model" if m["role"] == "assistant" else "user", parts=_to_parts(m["content"]))
        for m in messages
    ]
    config = types.GenerateContentConfig(system_instruction=system)
    last_error = None

    for model in [MODEL] + [m for m in FALLBACK_MODELS if m != MODEL]:
        for attempt in range(3):
            try:
                resp = client.models.generate_content(model=model, contents=contents, config=config)
                return (resp.text or "").strip()
            except Exception as e:
                last_error = e
                msg = str(e)
                if "503" in msg or "UNAVAILABLE" in msg or "500" in msg:
                    log.warning("%s saturado (intento %d/3)", model, attempt + 1)
                    time.sleep(2 * (attempt + 1))
                    continue
                if "404" in msg or "NOT_FOUND" in msg or "429" in msg or "RESOURCE_EXHAUSTED" in msg:
                    log.warning("%s no disponible o sin cuota: %s", model, msg[:120])
                    break  # pasar al siguiente modelo
                raise  # error distinto (clave inválida, petición mala...)

    msg = str(last_error)
    if "429" in msg or "RESOURCE_EXHAUSTED" in msg:
        raise RuntimeError("Se alcanzó el límite gratuito de la IA. Espera un minuto e inténtalo de nuevo.")
    raise RuntimeError("La IA está saturada ahora mismo. Inténtalo de nuevo en unos segundos.")


def parse_json_loose(text):
    """Extrae JSON aunque el modelo lo envuelva en ```json ... ```."""
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip(), flags=re.IGNORECASE)
    return json.loads(text)


def translate_text(text, direccion="auto"):
    if not text.strip():
        return ""
    system = BASE_RULES + "\n" + direction_rule(direccion) + "\nDevuelve ÚNICAMENTE la traducción."
    return ask_ai(system, [{"role": "user", "content": text}], max_tokens=8000)


def translate_batch(items, direccion="auto"):
    """Traduce una lista de textos (párrafos, celdas...) manteniendo el orden.
    Agrupa en lotes para no hacer una llamada por párrafo."""
    results = [""] * len(items)
    idx_with_text = [i for i, t in enumerate(items) if t.strip()]

    # Lotes por tamaño de texto y nº de elementos
    batches, current, size = [], [], 0
    for i in idx_with_text:
        if current and (size + len(items[i]) > 6000 or len(current) >= 60):
            batches.append(current)
            current, size = [], 0
        current.append(i)
        size += len(items[i])
    if current:
        batches.append(current)

    system = (
        BASE_RULES
        + "\n"
        + direction_rule(direccion)
        + "\nRecibirás un array JSON de textos. Devuelve ÚNICAMENTE un array JSON de strings "
        "con las traducciones, en el mismo orden y con exactamente la misma longitud. "
        "Cada texto es independiente. Sin comentarios ni bloques de código."
    )

    for batch in batches:
        payload = json.dumps([items[i] for i in batch], ensure_ascii=False)
        try:
            raw = ask_ai(system, [{"role": "user", "content": payload}], max_tokens=16000)
            translated = parse_json_loose(raw)
            if not isinstance(translated, list) or len(translated) != len(batch):
                raise ValueError("longitud distinta")
        except Exception as e:
            log.warning("Lote no válido (%s); traduzco elemento a elemento", e)
            translated = [translate_text(items[i], direccion) for i in batch]
        for i, t in zip(batch, translated):
            results[i] = str(t)
    return results


# ---------------------------------------------------------------------------
# Healthcheck
# ---------------------------------------------------------------------------
@app.route("/", methods=["GET"])
def healthcheck():
    return jsonify(status="ok", servicio="Parlem backend", modelo=MODEL)


# ---------------------------------------------------------------------------
# /chat
# ---------------------------------------------------------------------------
@app.route("/chat", methods=["POST"])
def chat():
    body = request.get_json(force=True, silent=True) or {}
    mensaje = (body.get("mensaje") or "").strip()
    historial = body.get("historial") or []
    direccion = clean_direction(body.get("direccion"))

    if not mensaje:
        return jsonify(error="Falta el mensaje"), 400

    # Solo aceptamos los últimos 20 turnos con formato válido
    messages = []
    for m in historial[-20:]:
        if m.get("role") in ("user", "assistant") and str(m.get("content", "")).strip():
            messages.append({"role": m["role"], "content": str(m["content"])})
    # La API exige que el primer mensaje sea del usuario y que alternen
    while messages and messages[0]["role"] != "user":
        messages.pop(0)
    merged = []
    for m in messages:
        if merged and merged[-1]["role"] == m["role"]:
            merged[-1]["content"] += "\n" + m["content"]
        else:
            merged.append(m)
    if merged and merged[-1]["role"] == "user":
        merged.pop()  # el mensaje actual se añade abajo
    merged.append({"role": "user", "content": mensaje})

    system = CHAT_SYSTEM
    if direccion != "auto":
        system += "\n" + direction_rule(direccion) + " (preferencia fijada por el usuario en la interfaz)."

    try:
        respuesta = ask_ai(system, merged)
        return jsonify(respuesta=respuesta)
    except Exception as e:
        log.exception("Error en /chat")
        return jsonify(error=str(e)), 502


# ---------------------------------------------------------------------------
# /traducir-texto
# ---------------------------------------------------------------------------
@app.route("/traducir-texto", methods=["POST"])
def traducir_texto():
    body = request.get_json(force=True, silent=True) or {}
    texto = body.get("texto", "")
    direccion = clean_direction(body.get("direccion"))
    try:
        return jsonify(traduccion=translate_text(texto, direccion))
    except Exception as e:
        log.exception("Error traduciendo texto")
        return jsonify(error=str(e)), 502


# ---------------------------------------------------------------------------
# /traducir-archivo  (Word, PowerPoint, PDF)
# ---------------------------------------------------------------------------
MIMETYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
}


@app.route("/traducir-archivo", methods=["POST"])
def traducir_archivo():
    if "file" not in request.files:
        return jsonify(error="Falta el archivo"), 400
    file = request.files["file"]
    direccion = clean_direction(request.form.get("direccion"))

    filename = file.filename or "archivo"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "doc":
        return jsonify(error="El formato .doc antiguo no se admite: guárdalo como .docx y vuelve a subirlo"), 400
    if ext not in MIMETYPES:
        return jsonify(error=f"Formato .{ext} no soportado (usa docx, pptx o pdf)"), 400

    data = file.read()
    try:
        if ext == "docx":
            out = translate_docx(data, direccion)
        elif ext == "pptx":
            out = translate_pptx(data, direccion)
        else:
            out = translate_pdf(data, direccion)
    except Exception as e:
        log.exception("Error traduciendo archivo")
        return jsonify(error=str(e)), 502

    out_name = filename.rsplit(".", 1)[0] + "_traducido." + ext
    return send_file(
        io.BytesIO(out),
        mimetype=MIMETYPES[ext],
        as_attachment=True,
        download_name=out_name,
    )


def replace_paragraph_text(paragraph, new_text):
    """Pone el texto traducido en el primer 'run' (conserva su formato)
    y vacía el resto."""
    runs = paragraph.runs
    if not runs:
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def translate_docx(data, direccion):
    from docx import Document

    doc = Document(io.BytesIO(data))

    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        for part in (section.header, section.footer):
            paragraphs.extend(part.paragraphs)

    # Evitar traducir dos veces celdas combinadas (mismo objeto XML)
    seen, unique = set(), []
    for p in paragraphs:
        if id(p._p) in seen:
            continue
        seen.add(id(p._p))
        if p.text.strip() and p.runs:
            unique.append(p)

    translated = translate_batch([p.text for p in unique], direccion)
    for p, t in zip(unique, translated):
        replace_paragraph_text(p, t)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def translate_pptx(data, direccion):
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    paragraphs = []

    def collect(shapes):
        for shape in shapes:
            if shape.shape_type == 6 and hasattr(shape, "shapes"):  # grupo
                collect(shape.shapes)
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                paragraphs.extend(shape.text_frame.paragraphs)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.text_frame.paragraphs)

    for slide in prs.slides:
        collect(slide.shapes)
        if slide.has_notes_slide:
            paragraphs.extend(slide.notes_slide.notes_text_frame.paragraphs)

    paragraphs = [p for p in paragraphs if p.runs and "".join(r.text for r in p.runs).strip()]
    texts = ["".join(r.text for r in p.runs) for p in paragraphs]
    translated = translate_batch(texts, direccion)
    for p, t in zip(paragraphs, translated):
        replace_paragraph_text(p, t)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def translate_pdf(data, direccion):
    """Extrae el texto, lo traduce y genera un PDF nuevo (no conserva el
    diseño exacto). Si el PDF es un escaneo sin texto, se lo pasa a la IA
    para que lo lea visualmente."""
    import pdfplumber
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    pages_text = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages_text.append(page.extract_text() or "")

    if any(t.strip() for t in pages_text):
        translated_pages = translate_batch(pages_text, direccion)
    else:
        # PDF escaneado: lectura nativa de la IA
        b64 = base64.standard_b64encode(data).decode()
        system = (
            BASE_RULES
            + "\n"
            + direction_rule(direccion)
            + "\nLee todo el documento y devuelve ÚNICAMENTE su texto traducido, "
            "conservando el orden y los párrafos."
        )
        content = [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": b64}},
            {"type": "text", "text": "Traduce este documento."},
        ]
        translated_pages = [ask_ai(system, [{"role": "user", "content": content}], max_tokens=16000)]

    styles = getSampleStyleSheet()
    story = []
    for i, page_text in enumerate(translated_pages):
        for line in page_text.split("\n"):
            if line.strip():
                story.append(Paragraph(escape(line), styles["Normal"]))
        if i < len(translated_pages) - 1:
            story.append(Spacer(1, 0.6 * cm))

    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4).build(story or [Paragraph("", styles["Normal"])])
    return out.getvalue()


# ---------------------------------------------------------------------------
# /traducir-imagen  (lectura del texto + traducción, todo con la IA)
# ---------------------------------------------------------------------------
IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}


@app.route("/traducir-imagen", methods=["POST"])
def traducir_imagen():
    if "file" not in request.files:
        return jsonify(error="Falta la imagen"), 400
    file = request.files["file"]
    direccion = clean_direction(request.form.get("direccion"))

    media_type = file.mimetype if file.mimetype in IMAGE_TYPES else None
    if media_type is None:
        ext = (file.filename or "").rsplit(".", 1)[-1].lower()
        media_type = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
                      "gif": "image/gif", "webp": "image/webp"}.get(ext)
    if media_type is None:
        return jsonify(error="Formato de imagen no soportado (usa png, jpg, webp o gif)"), 400

    b64 = base64.standard_b64encode(file.read()).decode()
    system = (
        BASE_RULES
        + "\n"
        + direction_rule(direccion)
        + '\nTe doy una imagen. Transcribe fielmente todo el texto que aparezca y tradúcelo. '
        'Devuelve ÚNICAMENTE un objeto JSON con esta forma: '
        '{"texto_original": "...", "traduccion": "..."}. '
        'Si no hay texto legible, devuelve ambos campos vacíos y añade "aviso": "No se detectó texto en la imagen."'
    )
    content = [
        {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
        {"type": "text", "text": "Extrae y traduce el texto de esta imagen."},
    ]

    try:
        raw = ask_ai(system, [{"role": "user", "content": content}], max_tokens=8000)
        try:
            return jsonify(parse_json_loose(raw))
        except Exception:
            # Si el modelo no devolvió JSON válido, entregamos el texto tal cual
            return jsonify(texto_original="", traduccion=raw)
    except Exception as e:
        log.exception("Error procesando imagen")
        return jsonify(error=str(e)), 502


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
